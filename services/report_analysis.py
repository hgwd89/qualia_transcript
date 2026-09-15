"""
分析用フラットデータ .xlsx / .csv 生成（QDAソフト・二次分析用）
"""
import csv
import io
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import config
from models.project import Project
from models.participant import Participant
from models.interview import Interview
from models.segment import Segment, UtteranceMapping
from models.generated_file import GeneratedFile
from models import db
from services.file_manager import (
    open_output_target_for_write,
    prepare_output_target,
    register_generated_file,
)
from services.ordinary_artifact_provenance import (
    begin_ordinary_artifact_source_snapshot,
    ordinary_artifact_generation_params,
)


def _build_rows(project_id: int) -> list[list]:
    """全発言のフラットレコードを返す"""
    project      = Project.query.get(project_id)
    interviews   = Interview.query.filter_by(project_id=project_id).all()

    # 属性キー一覧（全参加者の union）
    attr_keys = []
    for iv in interviews:
        if iv.participant:
            for a in iv.participant.attributes:
                if a.attribute_key not in attr_keys:
                    attr_keys.append(a.attribute_key)

    header = [
        "participant_code", "display_name",
        *attr_keys,
        "interview_date", "section_title", "question_code", "question_text",
        "is_key_question", "segment_seq", "start_sec", "end_sec",
        "speaker_role", "text", "mapped_by", "confidence", "is_unclassified",
    ]

    rows = [header]
    for iv in interviews:
        p         = iv.participant
        p_code    = p.participant_code if p else ""
        p_name    = p.display_name if p else ""
        attr_vals = {}
        if p:
            for a in p.attributes:
                attr_vals[a.attribute_key] = a.attribute_value

        for seg in iv.segments:
            if not seg.utterance_mappings:
                um_list = [None]
            else:
                um_list = seg.utterance_mappings

            for um in um_list:
                q    = um.question   if um else None
                sect = q.section     if q  else None

                rows.append([
                    p_code, p_name,
                    *[attr_vals.get(k, "") for k in attr_keys],
                    iv.interview_date.isoformat() if iv.interview_date else "",
                    sect.title        if sect else "",
                    q.question_code   if q    else "",
                    q.question_text   if q    else "",
                    "TRUE" if (q and q.is_key_question) else "FALSE",
                    seg.seq,
                    seg.start_sec or "",
                    seg.end_sec   or "",
                    seg.speaker_role,
                    seg.text,
                    um.mapped_by       if um else "",
                    um.confidence      if um else "",
                    um.is_unclassified if um else "",
                ])
    return rows


def generate_analysis_xlsx(project_id: int) -> GeneratedFile:
    begin_ordinary_artifact_source_snapshot(project_id)
    try:
        rows = _build_rows(project_id)
        project = db.session.get(Project, int(project_id))
        if project is None:
            raise ValueError("project が見つかりません")
        generation_params_json = ordinary_artifact_generation_params(
            "analysis",
            project_id,
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "発言データ"

        for r_idx, row in enumerate(rows, 1):
            for c_idx, val in enumerate(row, 1):
                ws.cell(r_idx, c_idx, val)

        # ヘッダー書式
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2E4057")

        ws.freeze_panes = "A2"

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"分析データ_{project.name}_{ts}.xlsx"
        target = prepare_output_target(project_id, filename)
        opened = open_output_target_for_write(target)
        try:
            wb.save(opened.stream)
        finally:
            opened.close()
        return register_generated_file(
            target,
            project_id=project_id,
            file_type="analysis",
            file_format="xlsx",
            generation_params_json=generation_params_json,
            existing_write_reservation=True,
        )
    except Exception:
        db.session.rollback()
        raise


def generate_analysis_csv(project_id: int) -> GeneratedFile:
    begin_ordinary_artifact_source_snapshot(project_id)
    try:
        rows = _build_rows(project_id)
        project = db.session.get(Project, int(project_id))
        if project is None:
            raise ValueError("project が見つかりません")
        generation_params_json = ordinary_artifact_generation_params(
            "analysis",
            project_id,
        )

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"分析データ_{project.name}_{ts}.csv"
        target = prepare_output_target(project_id, filename)

        opened = open_output_target_for_write(target)
        text_stream = io.TextIOWrapper(opened.stream, encoding="utf-8-sig", newline="")
        try:
            writer = csv.writer(text_stream)
            writer.writerows(rows)
            text_stream.flush()
        finally:
            try:
                text_stream.detach()
            except (OSError, ValueError):
                pass
            opened.close()

        return register_generated_file(
            target,
            project_id=project_id,
            file_type="analysis",
            file_format="csv",
            generation_params_json=generation_params_json,
            existing_write_reservation=True,
        )
    except Exception:
        db.session.rollback()
        raise
