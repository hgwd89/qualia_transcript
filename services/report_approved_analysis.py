"""Generate formal XLSX output from human-approved AIAnalysis rows only."""
import json
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text

import config
from models import db
from models.analysis import AIAnalysis
from models.generated_file import GeneratedFile
from models.project import Project
from services.analysis_source_provenance import (
    PROVENANCE_KEY,
    AnalysisSourceProvenanceError,
    require_current_analysis_source_provenance,
)
from services.approved_analysis_currentness import formal_analysis_state_sha256
from services.file_manager import (
    open_output_target_for_write,
    prepare_output_target,
    register_generated_file,
)
from services.formal_artifact_integrity import sha256_managed_generation


SUMMARY_HEADERS = [
    "analysis_id",
    "analysis_type",
    "title",
    "participant_code",
    "question_code",
    "summary_text",
    "implications",
    "unresolved",
    "review_status",
    "review_note",
    "reviewed_at",
    "model_used",
    "created_at",
]

EVIDENCE_HEADERS = [
    "analysis_id",
    "analysis_type",
    "finding_no",
    "point",
    "evidence_quote",
    "source_segment_ids",
    "participant_codes",
    "question_codes",
    "confidence",
]


def _parse_content(analysis: AIAnalysis) -> dict:
    try:
        content = json.loads(analysis.content_json or "{}")
    except (json.JSONDecodeError, TypeError) as exc:
        raise ValueError(f"AIAnalysis id={analysis.id} の content_json が不正です") from exc
    if not isinstance(content, dict):
        raise ValueError(f"AIAnalysis id={analysis.id} の content_json がobjectではありません")
    return content


def _participant_code(analysis: AIAnalysis) -> str:
    if analysis.interview and analysis.interview.participant:
        return analysis.interview.participant.participant_code or ""
    return ""


def _question_code(analysis: AIAnalysis) -> str:
    if analysis.question:
        return analysis.question.question_code or ""
    return ""


def _iso(value) -> str:
    return value.isoformat() if value else ""


def _join(values) -> str:
    if not values:
        return ""
    return ",".join(str(v) for v in values)


def _begin_formal_export_snapshot() -> None:
    """Serialize source validation with formal output registration."""
    if db.session.new or db.session.dirty or db.session.deleted:
        raise RuntimeError("formal analysis export requires a clean database session")
    db.session.rollback()
    if db.engine.dialect.name == "sqlite":
        db.session.execute(text("BEGIN IMMEDIATE"))


def _approved_rows(
    project_id: int,
) -> tuple[list[list], list[list], int, list[int], dict[str, str], dict[str, str]]:
    analyses = (
        AIAnalysis.query
        .filter_by(project_id=project_id, review_status="approved")
        .order_by(AIAnalysis.created_at.asc(), AIAnalysis.id.asc())
        .all()
    )
    if not analyses:
        raise ValueError("承認済みのAI分析がありません")

    summary_rows = [SUMMARY_HEADERS]
    evidence_rows = [EVIDENCE_HEADERS]
    finding_count = 0
    analysis_ids = []
    provenance_hashes: dict[str, str] = {}
    formal_state_hashes: dict[str, str] = {}

    for analysis in analyses:
        try:
            require_current_analysis_source_provenance(analysis)
        except AnalysisSourceProvenanceError as exc:
            raise ValueError(
                f"承認済みAIAnalysis id={analysis.id} の生成元入力が現在のcanonical dataと一致しません: {exc}"
            ) from exc

        content = _parse_content(analysis)
        provenance = content.get(PROVENANCE_KEY) or {}
        analysis_ids.append(int(analysis.id))
        provenance_hashes[str(int(analysis.id))] = str(provenance.get("sha256") or "")

        findings = content.get("findings") or []
        if not isinstance(findings, list):
            raise ValueError(f"AIAnalysis id={analysis.id} の findings が配列ではありません")

        summary_rows.append([
            analysis.id,
            analysis.analysis_type,
            analysis.title or "",
            _participant_code(analysis),
            _question_code(analysis),
            analysis.summary_text or "",
            content.get("implications", "") or "",
            content.get("unresolved", "") or "",
            analysis.review_status,
            analysis.review_note or "",
            _iso(analysis.reviewed_at),
            analysis.model_used or "",
            _iso(analysis.created_at),
        ])

        for index, finding in enumerate(findings, start=1):
            if not isinstance(finding, dict):
                raise ValueError(
                    f"AIAnalysis id={analysis.id} finding #{index} がobjectではありません"
                )
            source_ids = finding.get("source_segment_ids") or []
            if not source_ids:
                raise ValueError(
                    f"AIAnalysis id={analysis.id} finding #{index} に source_segment_ids がありません"
                )
            evidence_quote = str(finding.get("evidence_quote") or "").strip()
            if not evidence_quote:
                raise ValueError(
                    f"AIAnalysis id={analysis.id} finding #{index} に evidence_quote がありません"
                )

            evidence_rows.append([
                analysis.id,
                analysis.analysis_type,
                index,
                finding.get("point", "") or "",
                evidence_quote,
                _join(source_ids),
                _join(finding.get("participant_codes") or []),
                _join(finding.get("question_codes") or []),
                finding.get("confidence", "") or "",
            ])
            finding_count += 1

        formal_state_hashes[str(int(analysis.id))] = formal_analysis_state_sha256(analysis)

    return (
        summary_rows,
        evidence_rows,
        finding_count,
        analysis_ids,
        provenance_hashes,
        formal_state_hashes,
    )


def _style_sheet(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E4057")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for col_idx in range(1, ws.max_column + 1):
        max_len = 0
        for row_idx in range(1, min(ws.max_row, 100) + 1):
            value = ws.cell(row_idx, col_idx).value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 60)


def generate_approved_analysis_xlsx(project_id: int) -> GeneratedFile:
    _begin_formal_export_snapshot()
    try:
        project = db.session.get(Project, int(project_id))
        if not project:
            raise ValueError("project が見つかりません")

        (
            summary_rows,
            evidence_rows,
            finding_count,
            analysis_ids,
            provenance_hashes,
            formal_state_hashes,
        ) = _approved_rows(project_id)

        wb = Workbook()
        ws_summary = wb.active
        ws_summary.title = "承認済AI分析"
        for row in summary_rows:
            ws_summary.append(row)
        _style_sheet(ws_summary)

        ws_evidence = wb.create_sheet("根拠引用")
        for row in evidence_rows:
            ws_evidence.append(row)
        _style_sheet(ws_evidence)

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"承認済AI分析_{project.name}_{ts}.xlsx"
        target = prepare_output_target(project_id, filename)
        opened = open_output_target_for_write(target)
        try:
            wb.save(opened.stream)
        finally:
            opened.close()

        artifact_sha256 = sha256_managed_generation(
            config.OUTPUT_DIR,
            target.stored_path,
            target.written_stat,
        )
        params = {
            "approved_only": True,
            "analysis_count": len(summary_rows) - 1,
            "finding_count": finding_count,
            "analysis_ids": analysis_ids,
            "source_provenance_sha256": provenance_hashes,
            "formal_analysis_state_sha256": formal_state_hashes,
            "artifact_sha256": artifact_sha256,
        }
        return register_generated_file(
            target,
            project_id=project_id,
            file_type="approved_analysis",
            file_format="xlsx",
            generation_params_json=json.dumps(params, ensure_ascii=False),
            existing_write_reservation=True,
        )
    except Exception:
        db.session.rollback()
        raise
