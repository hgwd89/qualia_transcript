"""
整形シート .xlsx 生成。

行：質問項目、列：インタビュー、セル：該当発言テキスト。
同一参加者の複数回インタビューを潰さず、複数フローと未マッピング発言も保持する。
"""
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config
from models import db
from models.generated_file import GeneratedFile
from models.interview import Interview
from models.project import Project
from models.segment import Segment, UtteranceMapping
from models.speaker_assignment import SpeakerAssignment
from services.file_manager import (
    open_output_target_for_write,
    prepare_output_target,
    register_generated_file,
)
from services.generated_file_source_provenance import (
    capture_generated_file_source_provenance,
)


_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_KEY_FILL = PatternFill("solid", fgColor="FCE4D6")
_THIN = Side(style="thin", color="AAAAAA")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_FLAG_ORDER = ("favorite", "quote", "exclude", "needs_review")


def _segment_flag_map(seg) -> dict[str, bool]:
    flags = {f.flag_type for f in (seg.segment_flags or [])}
    return {name: (name in flags) for name in _FLAG_ORDER}


def _segment_flag_value_line(seg) -> str:
    flag_map = _segment_flag_map(seg)
    return ",".join(
        f"{name}={'true' if flag_map[name] else 'false'}"
        for name in _FLAG_ORDER
    )


def _interview_label(iv: Interview) -> str:
    p = iv.participant
    code = p.participant_code if p and p.participant_code else "participant未設定"
    name = p.display_name if p and p.display_name else ""
    date = iv.interview_date.isoformat() if iv.interview_date else "日付未設定"
    name_line = f" / {name}" if name else ""
    return f"{code}{name_line}\n{date}\nInterview ID={iv.id}"


def _is_unclassified(seg: Segment) -> bool:
    mappings = list(seg.utterance_mappings or [])
    if not mappings:
        return True
    return all(um.is_unclassified or um.question_id is None for um in mappings)


def _effective_role(seg: Segment, assignment_map: dict[str, SpeakerAssignment]) -> str:
    """Resolve the human speaker assignment before falling back to Segment role."""
    assignment = assignment_map.get(seg.speaker_label or "")
    if assignment and assignment.speaker_role:
        return assignment.speaker_role
    return seg.speaker_role or "unknown"


def generate_formatted_sheet(project_id: int) -> GeneratedFile:
    source_provenance = capture_generated_file_source_provenance(
        "formatted_sheet",
        int(project_id),
    )
    project = Project.query.get(project_id)
    if not project:
        raise ValueError("project が見つかりません")

    interviews = (
        Interview.query
        .filter_by(project_id=project_id)
        .order_by(Interview.interview_date.asc(), Interview.id.asc())
        .all()
    )
    interview_ids = [int(iv.id) for iv in interviews]
    assignment_maps: dict[int, dict[str, SpeakerAssignment]] = {
        int(iv.id): {} for iv in interviews
    }
    if interview_ids:
        assignments = (
            SpeakerAssignment.query
            .filter(SpeakerAssignment.interview_id.in_(interview_ids))
            .all()
        )
        for assignment in assignments:
            if assignment.speaker_label:
                assignment_maps[int(assignment.interview_id)][assignment.speaker_label] = assignment

    flows = sorted(project.interview_flows, key=lambda f: f.id or 0)
    if not flows:
        raise ValueError("インタビューフローが設定されていません")
    multiple_flows = len(flows) > 1

    wb = Workbook()
    ws = wb.active
    ws.title = "整形シート"

    base_headers = ("セクション", "質問コード", "質問テキスト")
    for idx, label in enumerate(base_headers, start=1):
        cell = ws.cell(1, idx, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, horizontal="center")

    col_offset = 4
    for i, iv in enumerate(interviews):
        text_col = col_offset + (i * 2)
        flag_col = text_col + 1
        label = _interview_label(iv)

        c_text = ws.cell(1, text_col, f"{label}\n発話")
        c_text.font = Font(bold=True, color="FFFFFF")
        c_text.fill = _HEADER_FILL
        c_text.alignment = Alignment(wrap_text=True, horizontal="center")

        c_flag = ws.cell(
            1,
            flag_col,
            f"{label}\nfavorite,quote,exclude,needs_review",
        )
        c_flag.font = Font(bold=True, color="FFFFFF")
        c_flag.fill = _HEADER_FILL
        c_flag.alignment = Alignment(wrap_text=True, horizontal="center")

    max_col = 3 + (len(interviews) * 2)
    row = 2
    for flow in flows:
        for section in flow.sections:
            section_label = (
                f"{flow.title} / {section.title}"
                if multiple_flows
                else section.title
            )
            for q in section.questions:
                ws.cell(row, 1, section_label)
                ws.cell(row, 2, q.question_code)
                ws.cell(row, 3, q.question_text).alignment = Alignment(wrap_text=True)

                if q.is_key_question:
                    for col in range(1, max_col + 1):
                        ws.cell(row, col).fill = _KEY_FILL

                for i, iv in enumerate(interviews):
                    candidate_mappings = (
                        UtteranceMapping.query
                        .filter_by(question_id=q.id)
                        .join(Segment, UtteranceMapping.segment_id == Segment.id)
                        .filter(Segment.interview_id == iv.id)
                        .order_by(Segment.seq.asc(), UtteranceMapping.id.asc())
                        .all()
                    )
                    assignment_map = assignment_maps.get(int(iv.id), {})
                    mappings = [
                        mapping
                        for mapping in candidate_mappings
                        if _effective_role(mapping.segment, assignment_map) == "respondent"
                    ]
                    text_col = col_offset + (i * 2)
                    flag_col = text_col + 1

                    texts = "\n".join(f"・{m.segment.text}" for m in mappings)
                    flags = "\n".join(
                        f"・{_segment_flag_value_line(m.segment)}"
                        for m in mappings
                    )

                    ws.cell(row, text_col, texts).alignment = Alignment(
                        wrap_text=True, vertical="top"
                    )
                    ws.cell(row, flag_col, flags).alignment = Alignment(
                        wrap_text=True, vertical="top"
                    )

                for col in range(1, max_col + 1):
                    ws.cell(row, col).border = _BORDER
                row += 1

    ws.column_dimensions["A"].width = 26 if multiple_flows else 18
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 40
    for i in range(len(interviews)):
        text_col = col_offset + (i * 2)
        flag_col = text_col + 1
        ws.column_dimensions[get_column_letter(text_col)].width = 40
        ws.column_dimensions[get_column_letter(flag_col)].width = 24
    ws.freeze_panes = "D2"

    ws2 = wb.create_sheet("未分類発言")
    ws2.append([
        "interview_id",
        "参加者",
        "実施日",
        "segment_id",
        "発言テキスト",
        "開始時刻",
        "speaker_label",
        "favorite",
        "quote",
        "exclude",
        "needs_review",
        "mapping_state",
    ])
    for cell in ws2[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL

    for iv in interviews:
        p = iv.participant
        code = p.participant_code if p and p.participant_code else "?"
        date = iv.interview_date.isoformat() if iv.interview_date else ""
        assignment_map = assignment_maps.get(int(iv.id), {})
        for seg in sorted(iv.segments, key=lambda s: (s.seq, s.id or 0)):
            if _effective_role(seg, assignment_map) != "respondent" or not _is_unclassified(seg):
                continue
            flag_map = _segment_flag_map(seg)
            mapping_state = "no_mapping" if not seg.utterance_mappings else "unclassified"
            ws2.append([
                iv.id,
                code,
                date,
                seg.id,
                seg.text,
                _fmt_time(seg.start_sec) if seg.start_sec is not None else "",
                seg.speaker_label or "",
                "true" if flag_map["favorite"] else "false",
                "true" if flag_map["quote"] else "false",
                "true" if flag_map["exclude"] else "false",
                "true" if flag_map["needs_review"] else "false",
                mapping_state,
            ])

    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions
    for col, width in {
        "A": 12, "B": 14, "C": 12, "D": 12, "E": 60, "F": 12,
        "G": 18, "H": 11, "I": 9, "J": 9, "K": 14, "L": 16,
    }.items():
        ws2.column_dimensions[col].width = width
    for row_cells in ws2.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"整形シート_{project.name}_{ts}.xlsx"
    target = prepare_output_target(project_id, filename)
    opened = open_output_target_for_write(target)
    try:
        wb.save(opened.stream)
    finally:
        opened.close()
    return register_generated_file(
        target,
        project_id=project_id,
        file_type="formatted_sheet",
        file_format="xlsx",
        source_provenance=source_provenance,
    )


def _fmt_time(sec: float) -> str:
    m, s = divmod(int(sec), 60)
    h, m = divmod(m, 60)
    return f"{h:02d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"
