"""
発言録 .docx 生成。

納品用の発言録は分析用マッピングに依存させず、Interview に属する全 Segment を
seq 順に欠落なく出力する。質問別の整理は formatted sheet 側の責務とする。
"""
import os
from datetime import datetime

from docx import Document
from docx.shared import Pt

import config
from models import db
from models.generated_file import GeneratedFile
from models.interview import Interview
from models.speaker_assignment import SpeakerAssignment
from services.file_manager import (
    open_output_target_for_write,
    prepare_output_target,
    register_generated_file,
)
from services.generated_file_source_provenance import (
    capture_generated_file_source_provenance,
)


ROLE_LABELS = {
    "moderator": "モデレーター",
    "interviewer": "モデレーター",
    "respondent": "参加者",
    "observer": "オブザーバー",
    "unknown": "話者未確定",
}


def _fmt_time(sec: float | None) -> str:
    if sec is None:
        return "--:--"
    m, s = divmod(int(sec), 60)
    h, m = divmod(m, 60)
    return f"{h:02d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def _has_quote_flag(seg) -> bool:
    return any(f.flag_type == "quote" for f in (seg.segment_flags or []))


def _speaker_name(interview: Interview, seg, assignment_map: dict) -> str:
    """人が確認しやすい話者名を返す。元の speaker_label も失わない。"""
    assignment = assignment_map.get(seg.speaker_label or "")
    role = (
        assignment.speaker_role
        if assignment and assignment.speaker_role
        else (seg.speaker_role or "unknown")
    )

    participant = None
    if assignment and assignment.participant:
        participant = assignment.participant
    elif seg.participant:
        participant = seg.participant
    elif role == "respondent":
        participant = interview.participant

    if participant:
        display = participant.display_name or participant.participant_code or "参加者"
    elif role in {"moderator", "interviewer"} and interview.interviewer_name:
        display = interview.interviewer_name
    else:
        display = ROLE_LABELS.get(role, role or "話者")

    label = seg.speaker_label or ""
    return f"{display} [{label}]" if label and label not in display else display


def generate_verbatim(interview_id: int) -> GeneratedFile:
    # Resolve only the owning project ID before capturing the canonical source
    # snapshot.  Do not load the mutable Interview object first; otherwise the
    # SQLAlchemy identity map could retain pre-snapshot values used by the export.
    project_id = (
        db.session.query(Interview.project_id)
        .filter(Interview.id == int(interview_id))
        .scalar()
    )
    if project_id is None:
        raise ValueError("interview が見つかりません")
    source_provenance = capture_generated_file_source_provenance(
        "verbatim",
        int(project_id),
        interview_id=int(interview_id),
    )

    interview = Interview.query.get(interview_id)
    if not interview:
        raise ValueError("interview が見つかりません")

    participant = interview.participant
    project = interview.project
    segments = sorted(
        interview.segments,
        key=lambda s: (
            s.seq if s.seq is not None else 10**9,
            s.start_sec if s.start_sec is not None else 10**12,
            s.id or 0,
        ),
    )

    assignments = SpeakerAssignment.query.filter_by(interview_id=interview.id).all()
    assignment_map = {a.speaker_label: a for a in assignments if a.speaker_label}

    doc = Document()

    title = doc.add_heading(level=1)
    run = title.add_run(
        f"発言録｜{participant.display_name if participant else '参加者未設定'}"
        f"｜{interview.interview_date or '日付未設定'}"
    )
    run.font.size = Pt(14)

    doc.add_paragraph(f"プロジェクト：{project.name}")
    if project.client:
        doc.add_paragraph(f"クライアント：{project.client}")
    if participant and participant.participant_code:
        doc.add_paragraph(f"参加者ID：{participant.participant_code}")
    if interview.interviewer_name:
        doc.add_paragraph(f"モデレーター：{interview.interviewer_name}")
    doc.add_paragraph(f"収録Segment数：{len(segments)}")
    doc.add_paragraph()

    doc.add_heading("逐語発言録（時系列）", level=2)

    for seg in segments:
        row = doc.add_paragraph()
        time_str = f"[{_fmt_time(seg.start_sec)}–{_fmt_time(seg.end_sec)}]"
        speaker = _speaker_name(interview, seg, assignment_map)
        lead = row.add_run(f"{time_str} {speaker}：")
        lead.bold = True
        if _has_quote_flag(seg):
            quote_mark = row.add_run("★引用候補 ")
            quote_mark.bold = True
        row.add_run(seg.text)

    if not segments:
        doc.add_paragraph("（発言Segmentがありません）")

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"発言録_{participant.participant_code if participant else 'unknown'}_{ts}.docx"
    target = prepare_output_target(interview.project_id, filename)
    opened = open_output_target_for_write(target)
    try:
        doc.save(opened.stream)
    finally:
        opened.close()
    return register_generated_file(
        target,
        project_id=project.id,
        interview_id=interview_id,
        file_type="verbatim",
        file_format="docx",
        source_provenance=source_provenance,
    )