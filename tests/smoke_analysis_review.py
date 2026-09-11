import json
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook


SEGMENT_TEXT = "保湿すると肌が落ち着いて安心します。"


def print_result(name: str, ok: bool, detail: str = "") -> bool:
    status = "PASS" if ok else "FAIL"
    suffix = f": {detail}" if detail else ""
    print(f"[{status}] {name}{suffix}")
    return ok


def add_failure(failures: int, name: str, ok: bool, detail: str = "") -> int:
    return failures + (0 if print_result(name, ok, detail) else 1)


def run_git(repo_root: Path, *args: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        ["git", *args],
        cwd=str(repo_root),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )


def repo_state(repo_root: Path) -> dict[str, bool]:
    return {
        "instance": (repo_root / "instance").exists(),
        "uploads": (repo_root / "uploads").exists(),
        "outputs": (repo_root / "outputs").exists(),
    }


def create_fixture(
    db,
    Project,
    Participant,
    Interview,
    InterviewFlow,
    InterviewFlowSection,
    InterviewFlowQuestion,
    Segment,
    UtteranceMapping,
    AIAnalysis,
):
    project = Project(name="Analysis Review Smoke", client="Smoke Client")
    db.session.add(project)
    db.session.flush()

    participant = Participant(
        project_id=project.id,
        participant_code="P01",
        display_name="Smoke Participant",
    )
    db.session.add(participant)
    db.session.flush()

    flow = InterviewFlow(project_id=project.id, title="Smoke Flow")
    db.session.add(flow)
    db.session.flush()

    section = InterviewFlowSection(flow_id=flow.id, title="Smoke Section", seq=1)
    db.session.add(section)
    db.session.flush()

    question = InterviewFlowQuestion(
        section_id=section.id,
        question_code="Q1",
        question_text="使った時の気持ちを教えてください",
        question_type="open",
        is_key_question=True,
        seq=1,
    )
    db.session.add(question)
    db.session.flush()

    interview = Interview(
        project_id=project.id,
        participant_id=participant.id,
        flow_id=flow.id,
        status="analyzed",
    )
    db.session.add(interview)
    db.session.flush()

    segment = Segment(
        interview_id=interview.id,
        participant_id=participant.id,
        speaker_label="RESP_A",
        speaker_role="respondent",
        start_sec=1.0,
        end_sec=4.0,
        text=SEGMENT_TEXT,
        seq=1,
    )
    db.session.add(segment)
    db.session.flush()

    db.session.add(
        UtteranceMapping(
            segment_id=segment.id,
            question_id=question.id,
            mapped_by="manual",
            confidence=1.0,
            is_unclassified=False,
        )
    )

    good_content = {
        "findings": [
            {
                "point": "保湿によって心理的な安心感も得ている",
                "evidence_quote": SEGMENT_TEXT,
                "participant_codes": ["P01"],
                "question_codes": ["Q1"],
                "confidence": "high",
            }
        ],
        "implications": "機能価値だけでなく安心感も訴求余地がある",
        "unresolved": "他参加者でも確認が必要",
    }
    bad_content = {
        "findings": [
            {
                "point": "原文に存在しない主張",
                "evidence_quote": "この発言は原文には存在しません。",
                "participant_codes": ["P01"],
                "question_codes": ["Q1"],
                "confidence": "low",
            }
        ],
        "implications": "",
        "unresolved": "",
    }

    good = AIAnalysis(
        project_id=project.id,
        interview_id=interview.id,
        question_id=question.id,
        analysis_type="per_question",
        title="Q1 考察",
        summary_text="安心感が重要",
        content_json=json.dumps(good_content, ensure_ascii=False),
        model_used="smoke-model",
    )
    bad = AIAnalysis(
        project_id=project.id,
        interview_id=interview.id,
        question_id=question.id,
        analysis_type="per_question",
        title="Q1 不一致考察",
        summary_text="根拠不一致",
        content_json=json.dumps(bad_content, ensure_ascii=False),
        model_used="smoke-model",
    )
    db.session.add_all([good, bad])
    db.session.commit()

    return {
        "project_id": project.id,
        "segment_id": segment.id,
        "good_analysis_id": good.id,
        "bad_analysis_id": bad.id,
    }


def main() -> int:
    repo_root = Path(__file__).resolve().parents[1]
    if str(repo_root) not in sys.path:
        sys.path.insert(0, str(repo_root))

    failures = 0
    before_repo_state = repo_state(repo_root)

    try:
        import config

        original_config = {
            "DATABASE_URI": config.DATABASE_URI,
            "UPLOAD_DIR": config.UPLOAD_DIR,
            "OUTPUT_DIR": config.OUTPUT_DIR,
        }
    except Exception as e:
        print_result("config import", False, f"{type(e).__name__}: {e}")
        return 1

    with tempfile.TemporaryDirectory(prefix="qualia_analysis_review_") as tmp:
        tmp_dir = Path(tmp)
        try:
            config.DATABASE_URI = f"sqlite:///{(tmp_dir / 'analysis_review.db').as_posix()}"
            config.UPLOAD_DIR = str(tmp_dir / "uploads")
            config.OUTPUT_DIR = str(tmp_dir / "outputs")

            from app import create_app
            from models import db
            from models.analysis import AIAnalysis
            from models.interview import Interview
            from models.interview_flow import InterviewFlow, InterviewFlowQuestion, InterviewFlowSection
            from models.participant import Participant
            from models.project import Project
            from models.segment import Segment, UtteranceMapping
            from services.report_approved_analysis import generate_approved_analysis_xlsx

            app = create_app()
            client = app.test_client()

            with app.app_context():
                ids = create_fixture(
                    db,
                    Project,
                    Participant,
                    Interview,
                    InterviewFlow,
                    InterviewFlowSection,
                    InterviewFlowQuestion,
                    Segment,
                    UtteranceMapping,
                    AIAnalysis,
                )
                baseline_text = db.session.get(Segment, ids["segment_id"]).text

            review_page = client.get(f"/projects/{ids['project_id']}/analysis/review")
            failures = add_failure(
                failures,
                "analysis review page renders",
                review_page.status_code == 200 and "AI分析承認".encode("utf-8") in review_page.data,
                f"status={review_page.status_code}",
            )

            approve_good = client.post(
                f"/api/projects/{ids['project_id']}/analysis/{ids['good_analysis_id']}/review",
                json={"status": "approved", "note": "smoke approved"},
            )
            good_json = approve_good.get_json(silent=True) or {}
            failures = add_failure(
                failures,
                "matching evidence can be approved",
                approve_good.status_code == 200
                and good_json.get("ok") is True
                and good_json.get("review_status") == "approved",
                f"status={approve_good.status_code}",
            )

            with app.app_context():
                good = db.session.get(AIAnalysis, ids["good_analysis_id"])
                good_content = json.loads(good.content_json)
                source_ids = good_content["findings"][0].get("source_segment_ids") or []
                failures = add_failure(
                    failures,
                    "approval stores source_segment_ids",
                    source_ids == [ids["segment_id"]],
                    str(source_ids),
                )

            approve_bad = client.post(
                f"/api/projects/{ids['project_id']}/analysis/{ids['bad_analysis_id']}/review",
                json={"status": "approved"},
            )
            bad_json = approve_bad.get_json(silent=True) or {}
            failures = add_failure(
                failures,
                "unmatched evidence is blocked from approval",
                approve_bad.status_code == 409 and bool(bad_json.get("unresolved")),
                f"status={approve_bad.status_code}",
            )

            with app.app_context():
                bad = db.session.get(AIAnalysis, ids["bad_analysis_id"])
                failures = add_failure(
                    failures,
                    "blocked analysis remains draft",
                    (bad.review_status or "draft") == "draft",
                    str(bad.review_status),
                )

            reject_bad = client.post(
                f"/api/projects/{ids['project_id']}/analysis/{ids['bad_analysis_id']}/review",
                json={"status": "rejected", "note": "evidence mismatch"},
            )
            reject_json = reject_bad.get_json(silent=True) or {}
            failures = add_failure(
                failures,
                "analysis can be rejected",
                reject_bad.status_code == 200 and reject_json.get("review_status") == "rejected",
                f"status={reject_bad.status_code}",
            )

            with app.app_context():
                gf = generate_approved_analysis_xlsx(ids["project_id"])
                output_path = Path(config.OUTPUT_DIR) / gf.stored_path
                failures = add_failure(
                    failures,
                    "approved analysis workbook generated in temp output",
                    output_path.is_file() and output_path.is_relative_to(tmp_dir),
                    str(output_path),
                )

                wb = load_workbook(output_path, data_only=True)
                failures = add_failure(
                    failures,
                    "approved workbook sheets",
                    set(wb.sheetnames) == {"承認済AI分析", "根拠引用"},
                    ", ".join(wb.sheetnames),
                )

                ws_summary = wb["承認済AI分析"]
                summary_ids = [ws_summary.cell(row, 1).value for row in range(2, ws_summary.max_row + 1)]
                failures = add_failure(
                    failures,
                    "formal output contains approved analysis only",
                    summary_ids == [ids["good_analysis_id"]],
                    str(summary_ids),
                )

                ws_evidence = wb["根拠引用"]
                evidence_headers = [c.value for c in ws_evidence[1]]
                source_col = evidence_headers.index("source_segment_ids") + 1
                evidence_source_ids = str(ws_evidence.cell(2, source_col).value or "")
                failures = add_failure(
                    failures,
                    "formal output preserves source_segment_ids",
                    evidence_source_ids == str(ids["segment_id"]),
                    evidence_source_ids,
                )

                segment = db.session.get(Segment, ids["segment_id"])
                failures = add_failure(
                    failures,
                    "Segment.text unchanged",
                    segment.text == baseline_text == SEGMENT_TEXT,
                )

                db.session.remove()
                db.engine.dispose()
        except Exception as e:
            failures = add_failure(
                failures,
                "analysis review smoke",
                False,
                f"{type(e).__name__}: {e}",
            )
        finally:
            config.DATABASE_URI = original_config["DATABASE_URI"]
            config.UPLOAD_DIR = original_config["UPLOAD_DIR"]
            config.OUTPUT_DIR = original_config["OUTPUT_DIR"]

    after_repo_state = repo_state(repo_root)
    failures = add_failure(
        failures,
        "repo instance/uploads/outputs state unchanged",
        before_repo_state == after_repo_state,
        f"before={before_repo_state}, after={after_repo_state}",
    )

    status_proc = run_git(repo_root, "status", "--short")
    if status_proc.returncode == 0:
        out = status_proc.stdout.strip()
        print("git status --short:")
        print(out if out else "(clean)")
    else:
        failures = add_failure(
            failures,
            "git status --short",
            False,
            (status_proc.stderr or status_proc.stdout or "unknown git error").strip(),
        )

    if failures == 0:
        print("\nSummary: PASS")
        return 0

    print(f"\nSummary: FAIL ({failures} checks failed)")
    return 1


if __name__ == "__main__":
    sys.exit(main())
