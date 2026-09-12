import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook


def check(name: str, ok: bool, detail: str = "") -> int:
    status = "PASS" if ok else "FAIL"
    suffix = f": {detail}" if detail else ""
    print(f"[{status}] {name}{suffix}")
    return 0 if ok else 1


def main() -> int:
    failures = 0
    repo_root = Path(__file__).resolve().parents[1]
    if str(repo_root) not in sys.path:
        sys.path.insert(0, str(repo_root))

    import config

    original = {
        "DATABASE_URI": config.DATABASE_URI,
        "UPLOAD_DIR": config.UPLOAD_DIR,
        "OUTPUT_DIR": config.OUTPUT_DIR,
    }
    app = None

    with tempfile.TemporaryDirectory(prefix="qualia_formatted_assignment_") as tmp:
        root = Path(tmp)
        config.DATABASE_URI = f"sqlite:///{(root / 'formatted.db').as_posix()}"
        config.UPLOAD_DIR = str(root / "uploads")
        config.OUTPUT_DIR = str(root / "outputs")
        try:
            from app import create_app
            from models import db
            from models.interview import Interview
            from models.interview_flow import (
                InterviewFlow,
                InterviewFlowQuestion,
                InterviewFlowSection,
            )
            from models.participant import Participant
            from models.project import Project
            from models.segment import Segment, UtteranceMapping
            from models.speaker_assignment import SpeakerAssignment
            from services.report_formatted import generate_formatted_sheet

            app = create_app()
            app.config["TESTING"] = True

            with app.app_context():
                project = Project(name="Formatted assignment smoke")
                db.session.add(project)
                db.session.flush()

                participant = Participant(
                    project_id=project.id,
                    participant_code="P01",
                    display_name="Participant 1",
                )
                flow = InterviewFlow(project_id=project.id, title="Flow")
                db.session.add_all([participant, flow])
                db.session.flush()

                section = InterviewFlowSection(flow_id=flow.id, title="Section", seq=1)
                db.session.add(section)
                db.session.flush()
                question = InterviewFlowQuestion(
                    section_id=section.id,
                    question_code="Q1",
                    question_text="Question 1",
                    seq=1,
                )
                db.session.add(question)
                db.session.flush()

                interview = Interview(
                    project_id=project.id,
                    participant_id=participant.id,
                    flow_id=flow.id,
                    status="mapped",
                )
                db.session.add(interview)
                db.session.flush()

                mapped_segment = Segment(
                    interview_id=interview.id,
                    participant_id=None,
                    seq=1,
                    speaker_label="SPEAKER_01",
                    speaker_role="unknown",
                    text="mapped respondent via speaker assignment",
                )
                unmapped_segment = Segment(
                    interview_id=interview.id,
                    participant_id=None,
                    seq=2,
                    speaker_label="SPEAKER_01",
                    speaker_role="unknown",
                    text="unmapped respondent via speaker assignment",
                )
                moderator_segment = Segment(
                    interview_id=interview.id,
                    participant_id=None,
                    seq=3,
                    speaker_label="SPEAKER_00",
                    speaker_role="unknown",
                    text="moderator should not appear",
                )
                db.session.add_all([mapped_segment, unmapped_segment, moderator_segment])
                db.session.flush()

                db.session.add(UtteranceMapping(
                    segment_id=mapped_segment.id,
                    question_id=question.id,
                    mapped_by="manual",
                    confidence=1.0,
                    is_unclassified=False,
                ))
                db.session.add_all([
                    SpeakerAssignment(
                        interview_id=interview.id,
                        speaker_label="SPEAKER_01",
                        speaker_role="respondent",
                        participant_id=participant.id,
                    ),
                    SpeakerAssignment(
                        interview_id=interview.id,
                        speaker_label="SPEAKER_00",
                        speaker_role="moderator",
                    ),
                ])
                db.session.commit()
                project_id = int(project.id)

                generated = generate_formatted_sheet(project_id)
                output_path = Path(config.OUTPUT_DIR) / generated.stored_path
                failures += check("formatted workbook created", output_path.is_file(), str(output_path))

                workbook = load_workbook(output_path, data_only=True)
                main_sheet = workbook["整形シート"]
                unclassified = workbook["未分類発言"]

                main_text = "\n".join(
                    str(cell.value or "")
                    for row in main_sheet.iter_rows()
                    for cell in row
                )
                unclassified_text = "\n".join(
                    str(cell.value or "")
                    for row in unclassified.iter_rows()
                    for cell in row
                )

                failures += check(
                    "mapped speaker-assigned respondent is retained",
                    "mapped respondent via speaker assignment" in main_text,
                )
                failures += check(
                    "unmapped speaker-assigned respondent is retained",
                    "unmapped respondent via speaker assignment" in unclassified_text,
                )
                failures += check(
                    "speaker-assigned moderator is excluded",
                    "moderator should not appear" not in main_text
                    and "moderator should not appear" not in unclassified_text,
                )

                db.session.remove()
                db.engine.dispose()
        except Exception as exc:
            failures += check(
                "formatted speaker-assignment smoke",
                False,
                f"{type(exc).__name__}: {exc}",
            )
        finally:
            if app is not None:
                try:
                    from models import db
                    with app.app_context():
                        db.session.remove()
                        db.engine.dispose()
                except Exception:
                    pass
            config.DATABASE_URI = original["DATABASE_URI"]
            config.UPLOAD_DIR = original["UPLOAD_DIR"]
            config.OUTPUT_DIR = original["OUTPUT_DIR"]

    if failures:
        print(f"\nSummary: FAIL ({failures} checks failed)")
        return 1
    print("\nSummary: PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
