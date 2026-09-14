import io
import sys
import tempfile
import zipfile
from pathlib import Path


def check(name: str, ok: bool, detail: str = "") -> int:
    status = "PASS" if ok else "FAIL"
    suffix = f": {detail}" if detail else ""
    print(f"[{status}] {name}{suffix}")
    return 0 if ok else 1


def has_code(items, code: str, *, file_id: int | None = None) -> bool:
    for item in items:
        if item.get("code") != code:
            continue
        if file_id is None:
            return True
        context = item.get("context") or {}
        if int(context.get("generated_file_id") or 0) == int(file_id):
            return True
    return False


def fake_zip_bytes() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("payload.txt", "this is a valid ZIP but not an OOXML package")
    return buffer.getvalue()


def main() -> int:
    failures = 0
    repo_root = Path(__file__).resolve().parents[1]
    scripts_dir = repo_root / "scripts"
    for entry in (repo_root, scripts_dir):
        if str(entry) not in sys.path:
            sys.path.insert(0, str(entry))

    import config
    from docx import Document
    from openpyxl import Workbook
    from services.readiness_validation import validate_generated_artifact

    original = {
        "DATABASE_URI": config.DATABASE_URI,
        "OUTPUT_DIR": config.OUTPUT_DIR,
        "UPLOAD_DIR": config.UPLOAD_DIR,
        "BACKUP_DIR": getattr(config, "BACKUP_DIR", None),
    }

    with tempfile.TemporaryDirectory(prefix="qualia_ooxml_readiness_") as tmp:
        root = Path(tmp)
        fake_xlsx = root / "fake.xlsx"
        fake_docx = root / "fake.docx"
        valid_xlsx = root / "valid.xlsx"
        valid_docx = root / "valid.docx"
        fake_bytes = fake_zip_bytes()
        fake_xlsx.write_bytes(fake_bytes)
        fake_docx.write_bytes(fake_bytes)

        workbook = Workbook()
        workbook.active["A1"] = "valid"
        workbook.save(valid_xlsx)
        workbook.close()

        document = Document()
        document.add_paragraph("valid")
        document.save(valid_docx)

        fake_xlsx_reason = validate_generated_artifact(fake_xlsx, "xlsx")
        fake_docx_reason = validate_generated_artifact(fake_docx, "docx")
        failures += check(
            "valid ZIP without SpreadsheetML package identity is rejected as XLSX",
            bool(fake_xlsx_reason),
            str(fake_xlsx_reason or ""),
        )
        failures += check(
            "valid ZIP without WordprocessingML package identity is rejected as DOCX",
            bool(fake_docx_reason),
            str(fake_docx_reason or ""),
        )
        failures += check(
            "openpyxl workbook passes OOXML package validation",
            validate_generated_artifact(valid_xlsx, "xlsx") is None,
            str(validate_generated_artifact(valid_xlsx, "xlsx") or ""),
        )
        failures += check(
            "python-docx document passes OOXML package validation",
            validate_generated_artifact(valid_docx, "docx") is None,
            str(validate_generated_artifact(valid_docx, "docx") or ""),
        )

        db_path = root / "readiness.db"
        output_dir = root / "outputs"
        backup_dir = root / "backups"
        config.DATABASE_URI = f"sqlite:///{db_path.as_posix()}"
        config.OUTPUT_DIR = str(output_dir)
        config.UPLOAD_DIR = str(root / "uploads")
        config.BACKUP_DIR = str(backup_dir)

        try:
            from app import create_app
            from audit_production_readiness_v2 import audit
            from models import db
            from models.project import Project
            from services.file_manager import (
                open_output_target_for_write,
                prepare_output_target,
                register_generated_file,
            )

            app = create_app()
            with app.app_context():
                project = Project(name="OOXML readiness smoke")
                db.session.add(project)
                db.session.commit()
                project_id = int(project.id)

                target = prepare_output_target(project_id, "fake-analysis.xlsx")
                opened = open_output_target_for_write(target)
                try:
                    opened.stream.write(fake_bytes)
                finally:
                    opened.close()
                generated = register_generated_file(
                    target,
                    project_id=project_id,
                    file_type="analysis",
                    file_format="xlsx",
                )
                generated_id = int(generated.id)
                db.session.remove()
                db.engine.dispose()

            report = audit(db_path, output_dir, backup_dir)
            integrity = report.get("info", {}).get("generated_file_byte_integrity", {})
            failures += check(
                "hash-bound fake XLSX still proves its registered bytes",
                integrity.get("verified") == 1
                and not has_code(
                    report.get("blockers", []),
                    "generated_file_byte_integrity_invalid",
                    file_id=generated_id,
                ),
                f"info={integrity} blockers={report.get('blockers', [])}",
            )
            failures += check(
                "readiness blocks a hash-valid ZIP that is not a real XLSX package",
                has_code(
                    report.get("blockers", []),
                    "generated_file_invalid",
                    file_id=generated_id,
                ),
                f"blockers={report.get('blockers', [])}",
            )
        finally:
            config.DATABASE_URI = original["DATABASE_URI"]
            config.OUTPUT_DIR = original["OUTPUT_DIR"]
            config.UPLOAD_DIR = original["UPLOAD_DIR"]
            if original["BACKUP_DIR"] is not None:
                config.BACKUP_DIR = original["BACKUP_DIR"]

    if failures:
        print(f"\nSummary: FAIL ({failures} checks failed)")
        return 1
    print("\nSummary: PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
