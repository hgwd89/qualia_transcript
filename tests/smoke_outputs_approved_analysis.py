import json
import sys
import tempfile
from pathlib import Path

from flask import Flask
from openpyxl import load_workbook


def print_result(name: str, ok: bool, detail: str = "") -> bool:
    status = "PASS" if ok else "FAIL"
    if detail:
        print(f"[{status}] {name}: {detail}")
    else:
        print(f"[{status}] {name}")
    return ok


def main() -> int:
    repo_root = Path(__file__).resolve().parents[1]
    if str(repo_root) not in sys.path:
        sys.path.insert(0, str(repo_root))

    failures = 0
    repo_state_before = {
        "instance": (repo_root / "instance").exists(),
        "uploads": (repo_root / "uploads").exists(),
        "outputs": (repo_root / "outputs").exists(),
    }

    try:
        import config
        from models import db
        from models.analysis import AIAnalysis
        from models.generated_file import GeneratedFile
        from models.interview import Interview, MediaFile, Transcription  # noqa: F401
        from models.interview_flow import InterviewFlow, InterviewFlowQuestion, InterviewFlowSection
        from models.participant import Participant, ParticipantAttribute  # noqa: F401
        from models.project import Project
        from models.quote_candidate import QuoteCandidate  # noqa: F401
        from models.review_item import ReviewItem  # noqa: F401
        from models.segment import Segment, UtteranceMapping
        from models.segment_flag import SegmentFlag  # noqa: F401
        from models.speaker_assignment import SpeakerAssignment  # noqa: F401
        from services.report_analysis import generate_analysis_xlsx
    except Exception as e:
        print_result("imports", False, f"{type(e).__name__}: {e}")
        return 1

    app = Flask(__name__)
    app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///:memory:"
    app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
    db.init_app(app)

    old_output_dir = config.OUTPUT_DIR
    try:
        with app.app_context(), tempfile.TemporaryDirectory() as tmp_output_dir:
            config.OUTPUT_DIR = tmp_output_dir
            db.create_all()

            project = Project(name="Approved Analysis Output Smoke")
            db.session.add(project)
            db.session.flush()

            participant = Participant(project_id=project.id, participant_code="P01", display_name="User1")
            db.session.add(participant)
            db.session.flush()

            flow = InterviewFlow(project_id=project.id, title="Flow")
            db.session.add(flow)
            db.session.flush()
            section = InterviewFlowSection(flow_id=flow.id, title="Section", seq=1)
            db.session.add(section)
            db.session.flush()
            question = InterviewFlowQuestion(section_id=section.id, question_code="Q1", question_text="Q1", seq=1)
            db.session.add(question)
            db.session.flush()
            interview = Interview(project_id=project.id, participant_id=participant.id, flow_id=flow.id)
            db.session.add(interview)
            db.session.flush()
            segment = Segment(
                interview_id=interview.id,
                participant_id=participant.id,
                speaker_label="SPEAKER_00",
                speaker_role="respondent",
                start_sec=1.0,
                end_sec=2.0,
                text="SEGMENT_TEXT_UNCHANGED",
                seq=1,
            )
            db.session.add(segment)
            db.session.flush()
            mapping = UtteranceMapping(segment_id=segment.id, question_id=question.id, confidence=0.9, confidence_level="high")
            db.session.add(mapping)

            approved_trace = AIAnalysis(
                project_id=project.id,
                interview_id=interview.id,
                question_id=question.id,
                analysis_type="per_question",
                title="APPROVED_TRACE_ANALYSIS",
                summary_text="APPROVED_TRACE_SUMMARY",
                content_json=json.dumps({
                    "source_segment_quotes": [{"segment_id": segment.id, "text": segment.text}],
                    "findings": [],
                }, ensure_ascii=False),
                source_segment_ids=json.dumps([segment.id]),
                quote_ids=json.dumps(["QT-TRACE-001"]),
                model_used="none",
                status="approved",
            )
            approved_missing_trace = AIAnalysis(
                project_id=project.id,
                analysis_type="per_question",
                title="APPROVED_MISSING_TRACE_ANALYSIS",
                summary_text="SHOULD_NOT_APPEAR",
                content_json=json.dumps({"findings": []}, ensure_ascii=False),
                model_used="none",
                status="approved",
            )
            draft_trace = AIAnalysis(
                project_id=project.id,
                analysis_type="per_question",
                title="DRAFT_TRACE_ANALYSIS",
                summary_text="SHOULD_NOT_APPEAR_DRAFT",
                content_json=json.dumps({
                    "source_segment_quotes": [{"segment_id": segment.id, "text": segment.text}],
                }, ensure_ascii=False),
                source_segment_ids=json.dumps([segment.id]),
                model_used="none",
                status="draft",
            )
            db.session.add_all([approved_trace, approved_missing_trace, draft_trace])
            db.session.commit()

            original_text = segment.text
            gf = generate_analysis_xlsx(project.id)
            output_path = Path(config.OUTPUT_DIR) / gf.stored_path
            wb = load_workbook(output_path)
            sheet_names = wb.sheetnames
            ws = wb["承認済み分析"]
            headers = [cell.value for cell in ws[1]]
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            flat_values = [str(value) for row in rows for value in row if value is not None]

            failures += 0 if print_result(
                "xlsx generated under temp output dir",
                output_path.is_file() and str(output_path).startswith(str(tmp_output_dir)),
                str(output_path),
            ) else 1
            failures += 0 if print_result(
                "GeneratedFile file_type unchanged",
                gf.file_type == "analysis" and gf.file_format == "xlsx",
                f"file_type={gf.file_type}, file_format={gf.file_format}",
            ) else 1
            failures += 0 if print_result(
                "approved analysis sheet exists",
                "承認済み分析" in sheet_names,
                ",".join(sheet_names),
            ) else 1
            failures += 0 if print_result(
                "approved analysis sheet headers",
                headers == [
                    "analysis_id",
                    "analysis_type",
                    "interview_id",
                    "question_id",
                    "title",
                    "summary_text",
                    "source_segment_ids",
                    "quote_ids",
                    "model_used",
                    "status",
                    "created_at",
                ],
                f"headers={headers}",
            ) else 1
            failures += 0 if print_result(
                "approved trace analysis is included",
                "APPROVED_TRACE_ANALYSIS" in flat_values and "APPROVED_TRACE_SUMMARY" in flat_values,
            ) else 1
            failures += 0 if print_result(
                "approved missing trace analysis is excluded",
                "APPROVED_MISSING_TRACE_ANALYSIS" not in flat_values and "SHOULD_NOT_APPEAR" not in flat_values,
            ) else 1
            failures += 0 if print_result(
                "draft trace analysis is excluded",
                "DRAFT_TRACE_ANALYSIS" not in flat_values and "SHOULD_NOT_APPEAR_DRAFT" not in flat_values,
            ) else 1
            failures += 0 if print_result(
                "trace ids are written",
                str(segment.id) in flat_values and "QT-TRACE-001" in flat_values,
                f"values={flat_values}",
            ) else 1
            failures += 0 if print_result(
                "Segment.text remains unchanged",
                db.session.get(Segment, segment.id).text == original_text,
            ) else 1
            db.session.remove()
            db.engine.dispose()
    finally:
        config.OUTPUT_DIR = old_output_dir

    repo_state_after = {
        "instance": (repo_root / "instance").exists(),
        "uploads": (repo_root / "uploads").exists(),
        "outputs": (repo_root / "outputs").exists(),
    }
    failures += 0 if print_result(
        "repo instance/uploads/outputs state unchanged",
        repo_state_before == repo_state_after,
        f"before={repo_state_before}, after={repo_state_after}",
    ) else 1

    if failures == 0:
        print("\nSummary: PASS")
        return 0
    print(f"\nSummary: FAIL ({failures} checks failed)")
    return 1


if __name__ == "__main__":
    sys.exit(main())
