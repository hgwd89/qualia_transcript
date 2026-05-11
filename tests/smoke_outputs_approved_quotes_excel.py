import sys
import tempfile
from pathlib import Path

from flask import Flask
from openpyxl import load_workbook


def print_result(name: str, ok: bool, detail: str = "") -> bool:
    status = "PASS" if ok else "FAIL"
    if detail:
        print(f"[{status}] {name}: {detail}")
    else:
        print(f"[{status}] {name}")
    return ok


def _sheet_rows(ws) -> list[tuple]:
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        if all((v is None or v == "") for v in row):
            continue
        rows.append(row)
    return rows


def main() -> int:
    repo_root = Path(__file__).resolve().parents[1]
    if str(repo_root) not in sys.path:
        sys.path.insert(0, str(repo_root))

    failures = 0

    try:
        import config
        from models import db
        from models.analysis import AIAnalysis  # noqa: F401 (mapper registry)
        from models.generated_file import GeneratedFile
        from models.interview import Interview, MediaFile, Transcription  # noqa: F401
        from models.interview_flow import InterviewFlow, InterviewFlowQuestion, InterviewFlowSection
        from models.participant import Participant, ParticipantAttribute  # noqa: F401
        from models.project import Project
        from models.quote_candidate import QuoteCandidate
        from models.review_item import ReviewItem  # noqa: F401
        from models.segment import Segment, UtteranceMapping
        from models.segment_flag import SegmentFlag
        from models.speaker_assignment import SpeakerAssignment  # noqa: F401
        from services.report_formatted import generate_formatted_sheet
    except Exception as e:
        print_result("imports", False, f"{type(e).__name__}: {e}")
        return 1

    app = Flask(__name__)
    app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///:memory:"
    app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
    db.init_app(app)

    with app.app_context(), tempfile.TemporaryDirectory() as tmp_output_dir:
        old_output_dir = config.OUTPUT_DIR
        config.OUTPUT_DIR = tmp_output_dir
        try:
            db.create_all()

            project = Project(name="Approved Quote Excel Smoke")
            db.session.add(project)
            db.session.flush()

            participant_1 = Participant(project_id=project.id, participant_code="P01", display_name="User1")
            participant_2 = Participant(project_id=project.id, participant_code="P02", display_name="User2")
            db.session.add_all([participant_1, participant_2])
            db.session.flush()

            flow = InterviewFlow(project_id=project.id, title="Flow")
            db.session.add(flow)
            db.session.flush()

            section = InterviewFlowSection(flow_id=flow.id, title="Section", seq=1)
            db.session.add(section)
            db.session.flush()

            question = InterviewFlowQuestion(section_id=section.id, question_code="Q1", question_text="Q1", seq=1)
            db.session.add(question)
            db.session.flush()

            interview_1 = Interview(project_id=project.id, participant_id=participant_1.id, flow_id=flow.id, status="mapped")
            interview_2 = Interview(project_id=project.id, participant_id=participant_2.id, flow_id=flow.id, status="mapped")
            db.session.add_all([interview_1, interview_2])
            db.session.flush()

            seg_approved_1 = Segment(
                interview_id=interview_1.id,
                participant_id=participant_1.id,
                speaker_label="SPEAKER_00",
                speaker_role="respondent",
                start_sec=1.0,
                end_sec=2.0,
                text="APPROVED_EXCEL_QUOTE_1",
                seq=1,
            )
            seg_candidate = Segment(
                interview_id=interview_1.id,
                participant_id=participant_1.id,
                speaker_label="SPEAKER_01",
                speaker_role="respondent",
                start_sec=2.0,
                end_sec=3.0,
                text="CANDIDATE_EXCEL_QUOTE",
                seq=2,
            )
            seg_rejected = Segment(
                interview_id=interview_1.id,
                participant_id=participant_1.id,
                speaker_label="SPEAKER_02",
                speaker_role="respondent",
                start_sec=3.0,
                end_sec=4.0,
                text="REJECTED_EXCEL_QUOTE",
                seq=3,
            )
            seg_approved_early = Segment(
                interview_id=interview_1.id,
                participant_id=participant_1.id,
                speaker_label="SPEAKER_04",
                speaker_role="respondent",
                start_sec=0.5,
                end_sec=0.9,
                text="APPROVED_EXCEL_QUOTE_EARLY",
                seq=5,
            )
            seg_approved_same_start_a = Segment(
                interview_id=interview_1.id,
                participant_id=participant_1.id,
                speaker_label="SPEAKER_05",
                speaker_role="respondent",
                start_sec=1.0,
                end_sec=1.2,
                text="APPROVED_EXCEL_QUOTE_SAME_A",
                seq=6,
            )
            seg_approved_same_start_z = Segment(
                interview_id=interview_1.id,
                participant_id=participant_1.id,
                speaker_label="SPEAKER_06",
                speaker_role="respondent",
                start_sec=1.0,
                end_sec=1.3,
                text="APPROVED_EXCEL_QUOTE_SAME_Z",
                seq=7,
            )
            seg_approved_none = Segment(
                interview_id=interview_1.id,
                participant_id=participant_1.id,
                speaker_label="SPEAKER_07",
                speaker_role="respondent",
                start_sec=None,
                end_sec=None,
                text="APPROVED_EXCEL_QUOTE_NONE",
                seq=8,
            )
            seg_flag_only = Segment(
                interview_id=interview_1.id,
                participant_id=participant_1.id,
                speaker_label="SPEAKER_03",
                speaker_role="respondent",
                start_sec=4.0,
                end_sec=5.0,
                text="FLAG_ONLY_SEGMENT_TEXT",
                seq=4,
            )
            seg_approved_2 = Segment(
                interview_id=interview_2.id,
                participant_id=participant_2.id,
                speaker_label="SPEAKER_10",
                speaker_role="respondent",
                start_sec=1.5,
                end_sec=2.5,
                text="APPROVED_EXCEL_QUOTE_2",
                seq=1,
            )
            db.session.add_all(
                [
                    seg_approved_1,
                    seg_candidate,
                    seg_rejected,
                    seg_approved_early,
                    seg_approved_same_start_a,
                    seg_approved_same_start_z,
                    seg_approved_none,
                    seg_flag_only,
                    seg_approved_2,
                ]
            )
            db.session.flush()

            baseline_texts = {
                seg_approved_1.id: seg_approved_1.text,
                seg_candidate.id: seg_candidate.text,
                seg_rejected.id: seg_rejected.text,
                seg_approved_early.id: seg_approved_early.text,
                seg_approved_same_start_a.id: seg_approved_same_start_a.text,
                seg_approved_same_start_z.id: seg_approved_same_start_z.text,
                seg_approved_none.id: seg_approved_none.text,
                seg_flag_only.id: seg_flag_only.text,
                seg_approved_2.id: seg_approved_2.text,
            }

            db.session.add_all(
                [
                    UtteranceMapping(
                        segment_id=seg_approved_1.id,
                        question_id=question.id,
                        mapped_by="ai",
                        confidence=0.95,
                        confidence_level="high",
                        is_unclassified=False,
                    ),
                    UtteranceMapping(
                        segment_id=seg_candidate.id,
                        question_id=question.id,
                        mapped_by="ai",
                        confidence=0.60,
                        confidence_level="medium",
                        is_unclassified=False,
                    ),
                    UtteranceMapping(
                        segment_id=seg_rejected.id,
                        question_id=question.id,
                        mapped_by="ai",
                        confidence=0.60,
                        confidence_level="medium",
                        is_unclassified=False,
                    ),
                    UtteranceMapping(
                        segment_id=seg_flag_only.id,
                        question_id=None,
                        mapped_by="ai",
                        confidence=0.30,
                        confidence_level="low",
                        is_unclassified=True,
                    ),
                    UtteranceMapping(
                        segment_id=seg_approved_2.id,
                        question_id=question.id,
                        mapped_by="ai",
                        confidence=0.90,
                        confidence_level="high",
                        is_unclassified=False,
                    ),
                    UtteranceMapping(
                        segment_id=seg_approved_early.id,
                        question_id=question.id,
                        mapped_by="ai",
                        confidence=0.90,
                        confidence_level="high",
                        is_unclassified=False,
                    ),
                    UtteranceMapping(
                        segment_id=seg_approved_same_start_a.id,
                        question_id=question.id,
                        mapped_by="ai",
                        confidence=0.90,
                        confidence_level="high",
                        is_unclassified=False,
                    ),
                    UtteranceMapping(
                        segment_id=seg_approved_same_start_z.id,
                        question_id=question.id,
                        mapped_by="ai",
                        confidence=0.90,
                        confidence_level="high",
                        is_unclassified=False,
                    ),
                    UtteranceMapping(
                        segment_id=seg_approved_none.id,
                        question_id=question.id,
                        mapped_by="ai",
                        confidence=0.90,
                        confidence_level="high",
                        is_unclassified=False,
                    ),
                ]
            )
            db.session.add(SegmentFlag(segment_id=seg_flag_only.id, flag_type="quote"))

            quote_approved_1 = QuoteCandidate(
                quote_id=f"QT-{interview_1.id}-{seg_approved_1.id}-APP",
                project_id=project.id,
                interview_id=interview_1.id,
                segment_id=seg_approved_1.id,
                participant_id=participant_1.id,
                question_id=question.id,
                start_sec=seg_approved_1.start_sec,
                end_sec=seg_approved_1.end_sec,
                quote_text=seg_approved_1.text,
                status="approved",
                source="human",
            )
            quote_candidate = QuoteCandidate(
                quote_id=f"QT-{interview_1.id}-{seg_candidate.id}-CAN",
                project_id=project.id,
                interview_id=interview_1.id,
                segment_id=seg_candidate.id,
                participant_id=participant_1.id,
                question_id=question.id,
                start_sec=seg_candidate.start_sec,
                end_sec=seg_candidate.end_sec,
                quote_text=seg_candidate.text,
                status="candidate",
                source="flag",
            )
            quote_rejected = QuoteCandidate(
                quote_id=f"QT-{interview_1.id}-{seg_rejected.id}-REJ",
                project_id=project.id,
                interview_id=interview_1.id,
                segment_id=seg_rejected.id,
                participant_id=participant_1.id,
                question_id=question.id,
                start_sec=seg_rejected.start_sec,
                end_sec=seg_rejected.end_sec,
                quote_text=seg_rejected.text,
                status="rejected",
                source="human",
            )
            quote_approved_2 = QuoteCandidate(
                quote_id=f"QT-{interview_2.id}-{seg_approved_2.id}-APP",
                project_id=project.id,
                interview_id=interview_2.id,
                segment_id=seg_approved_2.id,
                participant_id=participant_2.id,
                question_id=question.id,
                start_sec=seg_approved_2.start_sec,
                end_sec=seg_approved_2.end_sec,
                quote_text=seg_approved_2.text,
                status="approved",
                source="import",
            )
            quote_approved_early = QuoteCandidate(
                quote_id=f"QT-{interview_1.id}-{seg_approved_early.id}-EARLY",
                project_id=project.id,
                interview_id=interview_1.id,
                segment_id=seg_approved_early.id,
                participant_id=participant_1.id,
                question_id=question.id,
                start_sec=seg_approved_early.start_sec,
                end_sec=seg_approved_early.end_sec,
                quote_text=seg_approved_early.text,
                status="approved",
                source="human",
            )
            quote_approved_same_start_a = QuoteCandidate(
                quote_id=f"QT-{interview_1.id}-{seg_approved_same_start_a.id}-A",
                project_id=project.id,
                interview_id=interview_1.id,
                segment_id=seg_approved_same_start_a.id,
                participant_id=participant_1.id,
                question_id=question.id,
                start_sec=seg_approved_same_start_a.start_sec,
                end_sec=seg_approved_same_start_a.end_sec,
                quote_text=seg_approved_same_start_a.text,
                status="approved",
                source="human",
            )
            quote_approved_same_start_z = QuoteCandidate(
                quote_id=f"QT-{interview_1.id}-{seg_approved_same_start_z.id}-Z",
                project_id=project.id,
                interview_id=interview_1.id,
                segment_id=seg_approved_same_start_z.id,
                participant_id=participant_1.id,
                question_id=question.id,
                start_sec=seg_approved_same_start_z.start_sec,
                end_sec=seg_approved_same_start_z.end_sec,
                quote_text=seg_approved_same_start_z.text,
                status="approved",
                source="human",
            )
            quote_approved_none = QuoteCandidate(
                quote_id=f"QT-{interview_1.id}-{seg_approved_none.id}-NONE",
                project_id=project.id,
                interview_id=interview_1.id,
                segment_id=seg_approved_none.id,
                participant_id=participant_1.id,
                question_id=question.id,
                start_sec=seg_approved_none.start_sec,
                end_sec=seg_approved_none.end_sec,
                quote_text=seg_approved_none.text,
                status="approved",
                source="human",
            )
            db.session.add_all(
                [
                    quote_approved_1,
                    quote_candidate,
                    quote_rejected,
                    quote_approved_2,
                    quote_approved_early,
                    quote_approved_same_start_z,
                    quote_approved_same_start_a,
                    quote_approved_none,
                ]
            )
            db.session.commit()

            gf = generate_formatted_sheet(project.id)
            xlsx_path = Path(config.OUTPUT_DIR) / gf.stored_path

            failures += 0 if print_result(
                "xlsx generated under temp output dir",
                xlsx_path.is_file() and str(xlsx_path).startswith(str(Path(tmp_output_dir))),
                str(xlsx_path),
            ) else 1
            failures += 0 if print_result(
                "GeneratedFile file_type unchanged",
                gf.file_type == "formatted_sheet",
                f"file_type={gf.file_type}",
            ) else 1

            wb = load_workbook(str(xlsx_path), data_only=True)
            failures += 0 if print_result(
                "existing sheets are preserved",
                "整形シート" in wb.sheetnames and "未分類発言" in wb.sheetnames,
                ",".join(wb.sheetnames),
            ) else 1
            failures += 0 if print_result(
                "sheet order is preserved",
                wb.sheetnames == ["整形シート", "未分類発言", "正式引用（承認済み）"],
                ",".join(wb.sheetnames),
            ) else 1
            failures += 0 if print_result(
                "approved quote sheet exists",
                "正式引用（承認済み）" in wb.sheetnames,
                ",".join(wb.sheetnames),
            ) else 1

            ws_quotes = wb["正式引用（承認済み）"]
            headers = [cell.value for cell in ws_quotes[1]]
            expected_headers = [
                "interview_id",
                "participant_code",
                "quote_id",
                "segment_id",
                "question_id",
                "start_sec",
                "end_sec",
                "quote_text",
                "source",
                "status",
            ]
            failures += 0 if print_result(
                "approved quote sheet headers",
                headers == expected_headers,
                f"headers={headers}",
            ) else 1

            rows = _sheet_rows(ws_quotes)
            quote_ids = {row[2] for row in rows}
            statuses = {row[9] for row in rows}
            quote_texts = {row[7] for row in rows}
            row_order_keys = [(row[0], row[5], row[2]) for row in rows]

            failures += 0 if print_result(
                "approved quotes are included",
                {
                    quote_approved_1.quote_id,
                    quote_approved_2.quote_id,
                    quote_approved_early.quote_id,
                    quote_approved_same_start_a.quote_id,
                    quote_approved_same_start_z.quote_id,
                    quote_approved_none.quote_id,
                }.issubset(quote_ids),
                f"quote_ids={sorted(list(quote_ids))}",
            ) else 1
            failures += 0 if print_result(
                "candidate/rejected quotes are excluded",
                quote_candidate.quote_id not in quote_ids and quote_rejected.quote_id not in quote_ids,
                f"quote_ids={sorted(list(quote_ids))}",
            ) else 1
            failures += 0 if print_result(
                "SegmentFlag.quote-only segment is excluded",
                seg_flag_only.text not in quote_texts,
            ) else 1
            failures += 0 if print_result(
                "all rows are approved",
                statuses == {"approved"} if rows else True,
                f"statuses={sorted(list(statuses))}",
            ) else 1

            def _sort_key(key_row: tuple) -> tuple:
                interview_id, start_sec, quote_id = key_row
                return (
                    interview_id,
                    start_sec is None,
                    start_sec if start_sec is not None else float("inf"),
                    quote_id or "",
                )

            failures += 0 if print_result(
                "approved rows are sorted by interview_id/start_sec/quote_id",
                row_order_keys == sorted(row_order_keys, key=_sort_key),
                f"row_order_keys={row_order_keys}",
            ) else 1

            same_start_quote_ids = [
                row[2]
                for row in rows
                if row[0] == interview_1.id and row[5] == 1.0
            ]
            failures += 0 if print_result(
                "same start_sec rows are sorted by quote_id",
                len(same_start_quote_ids) >= 2 and same_start_quote_ids == sorted(same_start_quote_ids),
                f"same_start_quote_ids={same_start_quote_ids}",
            ) else 1

            post_texts = {s.id: s.text for s in Segment.query.order_by(Segment.id.asc()).all()}
            failures += 0 if print_result(
                "Segment.text remains unchanged",
                post_texts == baseline_texts,
            ) else 1
        finally:
            config.OUTPUT_DIR = old_output_dir

    if failures == 0:
        print("\nSummary: PASS")
        return 0
    print(f"\nSummary: FAIL ({failures} checks failed)")
    return 1


if __name__ == "__main__":
    sys.exit(main())
