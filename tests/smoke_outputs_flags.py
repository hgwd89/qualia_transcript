import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


def collect_doc_texts(doc: Document) -> list[str]:
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(p.text for p in cell.paragraphs)
    return texts


FLAG_TYPES = ("favorite", "quote", "exclude", "needs_review")
SEGMENT_TEXT = "出力フラグ確認用の発話です。"


def print_result(name: str, ok: bool, detail: str = "") -> bool:
    status = "PASS" if ok else "FAIL"
    if detail:
        print(f"[{status}] {name}: {detail}")
    else:
        print(f"[{status}] {name}")
    return ok


def run_git(repo_root: Path, *args: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        ["git", *args],
        cwd=str(repo_root),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )


def repo_state(repo_root: Path) -> dict[str, bool]:
    return {
        "instance": (repo_root / "instance").exists(),
        "uploads": (repo_root / "uploads").exists(),
        "outputs": (repo_root / "outputs").exists(),
    }


def create_fixture(
    db,
    Project,
    Participant,
    Interview,
    InterviewFlow,
    InterviewFlowSection,
    InterviewFlowQuestion,
    Segment,
    UtteranceMapping,
    SegmentFlag,
) -> dict[str, int]:
    project = Project(name="Output Flag Smoke Project", client="Smoke Client")
    db.session.add(project)
    db.session.flush()

    participant = Participant(
        project_id=project.id,
        participant_code="P01",
        display_name="Smoke Participant",
    )
    db.session.add(participant)
    db.session.flush()

    flow = InterviewFlow(project_id=project.id, title="Smoke Flow")
    db.session.add(flow)
    db.session.flush()

    section = InterviewFlowSection(flow_id=flow.id, title="Smoke Section", seq=1)
    db.session.add(section)
    db.session.flush()

    question = InterviewFlowQuestion(
        section_id=section.id,
        question_code="Q1",
        question_text="Smoke question",
        question_type="open",
        is_key_question=True,
        seq=1,
    )
    db.session.add(question)
    db.session.flush()

    interview = Interview(
        project_id=project.id,
        participant_id=participant.id,
        flow_id=flow.id,
        status="mapped",
    )
    db.session.add(interview)
    db.session.flush()

    segment = Segment(
        interview_id=interview.id,
        participant_id=participant.id,
        speaker_label="SPEAKER_00",
        speaker_role="respondent",
        start_sec=12.0,
        end_sec=18.0,
        text=SEGMENT_TEXT,
        seq=1,
    )
    db.session.add(segment)
    db.session.flush()

    db.session.add(
        UtteranceMapping(
            segment_id=segment.id,
            question_id=question.id,
            mapped_by="manual",
            confidence=1.0,
            is_unclassified=True,
        )
    )
    for flag_type in FLAG_TYPES:
        db.session.add(
            SegmentFlag(
                segment_id=segment.id,
                flag_type=flag_type,
                note=f"__smoke_outputs_flags_{flag_type}__",
            )
        )
    db.session.commit()
    return {
        "project_id": project.id,
        "interview_id": interview.id,
        "segment_id": segment.id,
    }


def main() -> int:
    repo_root = Path(__file__).resolve().parents[1]
    if str(repo_root) not in sys.path:
        sys.path.insert(0, str(repo_root))

    failures = 0
    before_repo_state = repo_state(repo_root)

    try:
        import config

        original_config = {
            "DATABASE_URI": config.DATABASE_URI,
            "UPLOAD_DIR": config.UPLOAD_DIR,
            "OUTPUT_DIR": config.OUTPUT_DIR,
        }
    except Exception as e:
        print_result("config import", False, f"{type(e).__name__}: {e}")
        return 1

    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        word_path = None
        excel_path = None
        final_text = None
        try:
            config.DATABASE_URI = f"sqlite:///{(tmp_dir / 'outputs_flags_smoke.db').as_posix()}"
            config.UPLOAD_DIR = str(tmp_dir / "uploads")
            config.OUTPUT_DIR = str(tmp_dir / "outputs")

            from app import create_app
            from models import db
            from models.interview import Interview
            from models.interview_flow import InterviewFlow, InterviewFlowQuestion, InterviewFlowSection
            from models.participant import Participant
            from models.project import Project
            from models.segment import Segment, UtteranceMapping
            from models.segment_flag import SegmentFlag
            from services.report_formatted import generate_formatted_sheet
            from services.report_verbatim import generate_verbatim

            app = create_app()
            failures += 0 if print_result(
                "temporary database configured",
                app.config.get("SQLALCHEMY_DATABASE_URI") == config.DATABASE_URI,
                app.config.get("SQLALCHEMY_DATABASE_URI", ""),
            ) else 1
            failures += 0 if print_result(
                "temporary output dir configured",
                str(tmp_dir) in config.OUTPUT_DIR,
                config.OUTPUT_DIR,
            ) else 1

            with app.app_context():
                ids = create_fixture(
                    db,
                    Project,
                    Participant,
                    Interview,
                    InterviewFlow,
                    InterviewFlowSection,
                    InterviewFlowQuestion,
                    Segment,
                    UtteranceMapping,
                    SegmentFlag,
                )
                segment_id = ids["segment_id"]
                project_id = ids["project_id"]
                interview_id = ids["interview_id"]
                baseline_text = db.session.get(Segment, segment_id).text
                failures += 0 if print_result(
                    "self-contained output fixture created",
                    baseline_text == SEGMENT_TEXT,
                    f"segment_id={segment_id}",
                ) else 1

                gf_word = generate_verbatim(interview_id)
                word_path = Path(config.OUTPUT_DIR) / gf_word.stored_path
                failures += 0 if print_result(
                    "generate_verbatim",
                    word_path.is_file() and word_path.is_relative_to(tmp_dir),
                    str(word_path),
                ) else 1

                gf_excel = generate_formatted_sheet(project_id)
                excel_path = Path(config.OUTPUT_DIR) / gf_excel.stored_path
                failures += 0 if print_result(
                    "generate_formatted_sheet",
                    excel_path.is_file() and excel_path.is_relative_to(tmp_dir),
                    str(excel_path),
                ) else 1

                final_text = db.session.get(Segment, segment_id).text
                db.session.remove()
                db.engine.dispose()

            if word_path is not None and word_path.is_file():
                doc = Document(str(word_path))
                texts = collect_doc_texts(doc)
                marker_found = any("★引用候補" in t for t in texts)
                target_text_found = any(SEGMENT_TEXT in t for t in texts)
                failures += 0 if print_result(
                    "word marker present",
                    marker_found,
                    f"path={word_path.name}",
                ) else 1
                failures += 0 if print_result(
                    "word target segment text present",
                    target_text_found,
                ) else 1
            else:
                failures += 0 if print_result("word output exists", False, str(word_path)) else 1

            if excel_path is not None and excel_path.is_file():
                wb = load_workbook(str(excel_path), data_only=True)
                ws_unclassified = wb["未分類発言"]
                headers = [str(c.value) if c.value is not None else "" for c in ws_unclassified[1]]
                header_ok = all(name in headers for name in FLAG_TYPES)
                failures += 0 if print_result(
                    "excel unclassified flag columns",
                    header_ok,
                    ",".join(headers),
                ) else 1

                flag_values_ok = False
                if header_ok and "発言テキスト" in headers:
                    text_col = headers.index("発言テキスト") + 1
                    idx = {name: headers.index(name) + 1 for name in FLAG_TYPES}
                    for row in range(2, ws_unclassified.max_row + 1):
                        if ws_unclassified.cell(row, text_col).value == SEGMENT_TEXT:
                            vals = {
                                k: str(ws_unclassified.cell(row, c).value or "").strip().lower()
                                for k, c in idx.items()
                            }
                            flag_values_ok = all(vals[k] == "true" for k in FLAG_TYPES)
                            break
                failures += 0 if print_result(
                    "excel flag values true for target segment",
                    flag_values_ok,
                ) else 1
            else:
                failures += 0 if print_result("excel output exists", False, str(excel_path)) else 1

            failures += 0 if print_result("segment text unchanged", final_text == SEGMENT_TEXT) else 1
        except Exception as e:
            failures += 0 if print_result("output flag smoke", False, f"{type(e).__name__}: {e}") else 1
        finally:
            config.DATABASE_URI = original_config["DATABASE_URI"]
            config.UPLOAD_DIR = original_config["UPLOAD_DIR"]
            config.OUTPUT_DIR = original_config["OUTPUT_DIR"]

    after_repo_state = repo_state(repo_root)
    failures += 0 if print_result(
        "repo instance/uploads/outputs state unchanged",
        before_repo_state == after_repo_state,
        f"before={before_repo_state}, after={after_repo_state}",
    ) else 1

    status_proc = run_git(repo_root, "status", "--short")
    if status_proc.returncode == 0:
        out = status_proc.stdout.strip()
        print("git status --short:")
        print(out if out else "(clean)")
    else:
        failures += 0 if print_result(
            "git status --short",
            False,
            (status_proc.stderr or status_proc.stdout or "unknown git error").strip(),
        ) else 1

    if failures == 0:
        print("\nSummary: PASS")
        return 0
    print(f"\nSummary: FAIL ({failures} checks failed)")
    return 1


if __name__ == "__main__":
    sys.exit(main())
