import sys
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


MOD_TEXT = "では最初に、普段のスキンケアについて教えてください。"
MAPPED_TEXT_1 = "毎晩、化粧水のあとに保湿クリームを使っています。"
UNMAPPED_TEXT = "冬は特に乾燥するので、量を増やします。"
MAPPED_TEXT_2 = "朝も保湿しますが、夜より少なめです。"
MULTI_FLOW_TEXT = "別フローでは購入時の比較行動について話しました。"


def check(name: str, ok: bool, detail: str = "") -> int:
    status = "PASS" if ok else "FAIL"
    suffix = f": {detail}" if detail else ""
    print(f"[{status}] {name}{suffix}")
    return 0 if ok else 1


def doc_text(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def main() -> int:
    repo_root = Path(__file__).resolve().parents[1]
    if str(repo_root) not in sys.path:
        sys.path.insert(0, str(repo_root))

    import config

    original = {
        "DATABASE_URI": config.DATABASE_URI,
        "UPLOAD_DIR": config.UPLOAD_DIR,
        "OUTPUT_DIR": config.OUTPUT_DIR,
    }
    failures = 0

    with tempfile.TemporaryDirectory(prefix="qualia_professional_integrity_") as tmp:
        tmp_dir = Path(tmp)
        try:
            config.DATABASE_URI = f"sqlite:///{(tmp_dir / 'professional.db').as_posix()}"
            config.UPLOAD_DIR = str(tmp_dir / "uploads")
            config.OUTPUT_DIR = str(tmp_dir / "outputs")

            from app import create_app
            from models import db
            from models.interview import Interview
            from models.interview_flow import InterviewFlow, InterviewFlowQuestion, InterviewFlowSection
            from models.participant import Participant
            from models.project import Project
            from models.segment import Segment, UtteranceMapping
            from models.speaker_assignment import SpeakerAssignment
            from services.report_formatted import generate_formatted_sheet
            from services.report_verbatim import generate_verbatim

            app = create_app()
            client = app.test_client()

            with app.app_context():
                project = Project(name="Professional Integrity Smoke", client="Smoke Client")
                other_project = Project(name="Other Project")
                db.session.add_all([project, other_project])
                db.session.flush()

                participant = Participant(
                    project_id=project.id,
                    participant_code="P01",
                    display_name="回答者A",
                )
                outsider = Participant(
                    project_id=other_project.id,
                    participant_code="OUT01",
                    display_name="別案件回答者",
                )
                db.session.add_all([participant, outsider])
                db.session.flush()

                flow = InterviewFlow(project_id=project.id, title="本番想定フロー")
                flow2 = InterviewFlow(project_id=project.id, title="追加フロー")
                db.session.add_all([flow, flow2])
                db.session.flush()

                section = InterviewFlowSection(flow_id=flow.id, title="基本行動", seq=1)
                section2 = InterviewFlowSection(flow_id=flow2.id, title="購入行動", seq=1)
                db.session.add_all([section, section2])
                db.session.flush()

                question = InterviewFlowQuestion(
                    section_id=section.id,
                    question_code="Q1",
                    question_text="普段のスキンケアを教えてください",
                    question_type="open",
                    is_key_question=True,
                    seq=1,
                )
                question2 = InterviewFlowQuestion(
                    section_id=section2.id,
                    question_code="QX1",
                    question_text="購入時に何を比較しますか",
                    question_type="open",
                    is_key_question=True,
                    seq=1,
                )
                db.session.add_all([question, question2])
                db.session.flush()

                iv1 = Interview(
                    project_id=project.id,
                    participant_id=participant.id,
                    flow_id=flow.id,
                    interview_date=date(2026, 9, 1),
                    interviewer_name="モデレーターA",
                    status="mapped",
                )
                iv2 = Interview(
                    project_id=project.id,
                    participant_id=participant.id,
                    flow_id=flow.id,
                    interview_date=date(2026, 9, 2),
                    interviewer_name="モデレーターA",
                    status="mapped",
                )
                iv3 = Interview(
                    project_id=project.id,
                    participant_id=participant.id,
                    flow_id=flow2.id,
                    interview_date=date(2026, 9, 3),
                    interviewer_name="モデレーターB",
                    status="mapped",
                )
                db.session.add_all([iv1, iv2, iv3])
                db.session.flush()

                s_mod = Segment(
                    interview_id=iv1.id,
                    speaker_label="SPEAKER_00",
                    speaker_role="moderator",
                    start_sec=0.0,
                    end_sec=5.0,
                    text=MOD_TEXT,
                    seq=1,
                )
                s_mapped1 = Segment(
                    interview_id=iv1.id,
                    participant_id=participant.id,
                    speaker_label="SPEAKER_01",
                    speaker_role="respondent",
                    start_sec=5.0,
                    end_sec=14.0,
                    text=MAPPED_TEXT_1,
                    seq=2,
                )
                s_unmapped = Segment(
                    interview_id=iv1.id,
                    participant_id=participant.id,
                    speaker_label="SPEAKER_01",
                    speaker_role="respondent",
                    start_sec=14.0,
                    end_sec=20.0,
                    text=UNMAPPED_TEXT,
                    seq=3,
                )
                s_mapped2 = Segment(
                    interview_id=iv2.id,
                    participant_id=participant.id,
                    speaker_label="SPEAKER_01",
                    speaker_role="respondent",
                    start_sec=2.0,
                    end_sec=8.0,
                    text=MAPPED_TEXT_2,
                    seq=1,
                )
                s_flow2 = Segment(
                    interview_id=iv3.id,
                    participant_id=participant.id,
                    speaker_label="SPEAKER_01",
                    speaker_role="respondent",
                    start_sec=1.0,
                    end_sec=7.0,
                    text=MULTI_FLOW_TEXT,
                    seq=1,
                )
                db.session.add_all([s_mod, s_mapped1, s_unmapped, s_mapped2, s_flow2])
                db.session.flush()

                db.session.add_all([
                    UtteranceMapping(
                        segment_id=s_mapped1.id,
                        question_id=question.id,
                        mapped_by="manual",
                        confidence=1.0,
                        is_unclassified=False,
                    ),
                    UtteranceMapping(
                        segment_id=s_mapped2.id,
                        question_id=question.id,
                        mapped_by="manual",
                        confidence=1.0,
                        is_unclassified=False,
                    ),
                    UtteranceMapping(
                        segment_id=s_flow2.id,
                        question_id=question2.id,
                        mapped_by="manual",
                        confidence=1.0,
                        is_unclassified=False,
                    ),
                    SpeakerAssignment(
                        interview_id=iv1.id,
                        speaker_label="SPEAKER_00",
                        speaker_role="moderator",
                    ),
                    SpeakerAssignment(
                        interview_id=iv1.id,
                        speaker_label="SPEAKER_01",
                        speaker_role="respondent",
                        participant_id=participant.id,
                    ),
                ])
                db.session.commit()

                ids = {
                    "project": project.id,
                    "iv1": iv1.id,
                    "iv2": iv2.id,
                    "iv3": iv3.id,
                    "s_mod": s_mod.id,
                    "s_mapped1": s_mapped1.id,
                    "s_unmapped": s_unmapped.id,
                    "participant": participant.id,
                    "outsider": outsider.id,
                }
                baseline = {
                    s.id: s.text
                    for s in (s_mod, s_mapped1, s_unmapped, s_mapped2, s_flow2)
                }

                gf_word = generate_verbatim(iv1.id)
                gf_xlsx = generate_formatted_sheet(project.id)
                word_path = Path(config.OUTPUT_DIR) / gf_word.stored_path
                xlsx_path = Path(config.OUTPUT_DIR) / gf_xlsx.stored_path

            text = doc_text(word_path)
            failures += check("verbatim contains moderator segment", MOD_TEXT in text)
            failures += check("verbatim contains mapped respondent segment", MAPPED_TEXT_1 in text)
            failures += check("verbatim contains zero-mapping respondent segment", UNMAPPED_TEXT in text)
            failures += check(
                "verbatim preserves chronological order",
                text.index(MOD_TEXT) < text.index(MAPPED_TEXT_1) < text.index(UNMAPPED_TEXT),
            )
            failures += check("verbatim identifies moderator", "モデレーターA" in text)
            failures += check("verbatim identifies respondent", "回答者A" in text)

            wb = load_workbook(xlsx_path, data_only=True)
            ws = wb["整形シート"]
            header_values = [str(c.value or "") for c in ws[1]]
            iv1_headers = [i for i, v in enumerate(header_values, start=1) if f"Interview ID={ids['iv1']}" in v]
            iv2_headers = [i for i, v in enumerate(header_values, start=1) if f"Interview ID={ids['iv2']}" in v]
            iv3_headers = [i for i, v in enumerate(header_values, start=1) if f"Interview ID={ids['iv3']}" in v]
            failures += check("formatted sheet keeps first interview", len(iv1_headers) == 2, str(iv1_headers))
            failures += check("formatted sheet keeps second interview for same participant", len(iv2_headers) == 2, str(iv2_headers))
            failures += check("formatted sheet keeps interview on second flow", len(iv3_headers) == 2, str(iv3_headers))

            q_row = None
            q2_row = None
            for row in range(2, ws.max_row + 1):
                if ws.cell(row, 2).value == "Q1":
                    q_row = row
                if ws.cell(row, 2).value == "QX1":
                    q2_row = row
            failures += check("formatted sheet Q1 row exists", q_row is not None)
            failures += check("formatted sheet second-flow question row exists", q2_row is not None)
            if q_row is not None and iv1_headers and iv2_headers:
                failures += check(
                    "first interview mapped text preserved",
                    MAPPED_TEXT_1 in str(ws.cell(q_row, iv1_headers[0]).value or ""),
                )
                failures += check(
                    "second interview mapped text preserved",
                    MAPPED_TEXT_2 in str(ws.cell(q_row, iv2_headers[0]).value or ""),
                )
            if q2_row is not None and iv3_headers:
                failures += check(
                    "second-flow mapped text preserved",
                    MULTI_FLOW_TEXT in str(ws.cell(q2_row, iv3_headers[0]).value or ""),
                )
                failures += check(
                    "second-flow section is disambiguated",
                    "追加フロー" in str(ws.cell(q2_row, 1).value or ""),
                )

            ws_un = wb["未分類発言"]
            headers = [c.value for c in ws_un[1]]
            seg_col = headers.index("segment_id") + 1
            state_col = headers.index("mapping_state") + 1
            text_col = headers.index("発言テキスト") + 1
            unmapped_rows = [
                r for r in range(2, ws_un.max_row + 1)
                if ws_un.cell(r, seg_col).value == ids["s_unmapped"]
            ]
            failures += check("zero-mapping segment appears in unclassified sheet", len(unmapped_rows) == 1)
            if unmapped_rows:
                r = unmapped_rows[0]
                failures += check("zero-mapping state is explicit", ws_un.cell(r, state_col).value == "no_mapping")
                failures += check("zero-mapping text preserved", ws_un.cell(r, text_col).value == UNMAPPED_TEXT)

            wrong_interview = client.post(
                f"/interviews/{ids['iv2']}/segments/{ids['s_mapped1']}/role",
                json={"speaker_role": "respondent", "participant_id": ids["participant"]},
            )
            failures += check("role update rejects cross-interview segment", wrong_interview.status_code == 400, str(wrong_interview.status_code))

            invalid_role = client.post(
                f"/interviews/{ids['iv1']}/segments/{ids['s_mapped1']}/role",
                json={"speaker_role": "admin", "participant_id": ids["participant"]},
            )
            failures += check("role update rejects invalid role", invalid_role.status_code == 400, str(invalid_role.status_code))

            outsider_resp = client.post(
                f"/interviews/{ids['iv1']}/segments/{ids['s_mapped1']}/role",
                json={"speaker_role": "respondent", "participant_id": ids["outsider"]},
            )
            failures += check("role update rejects participant from another project", outsider_resp.status_code == 400, str(outsider_resp.status_code))

            valid_resp = client.post(
                f"/interviews/{ids['iv1']}/segments/{ids['s_mapped1']}/role",
                json={"speaker_role": "respondent", "participant_id": ids["participant"]},
            )
            failures += check("role update accepts valid scoped update", valid_resp.status_code == 200, str(valid_resp.status_code))

            with app.app_context():
                from models.segment import Segment
                for segment_id, expected in baseline.items():
                    actual = db.session.get(Segment, segment_id).text
                    failures += check(f"Segment.text unchanged id={segment_id}", actual == expected)
                db.session.remove()
                db.engine.dispose()

        except Exception as exc:
            failures += check("professional integrity smoke", False, f"{type(exc).__name__}: {exc}")
            try:
                from models import db
                db.session.remove()
                db.engine.dispose()
            except Exception:
                pass
        finally:
            config.DATABASE_URI = original["DATABASE_URI"]
            config.UPLOAD_DIR = original["UPLOAD_DIR"]
            config.OUTPUT_DIR = original["OUTPUT_DIR"]

    if failures:
        print(f"\nSummary: FAIL ({failures} checks failed)")
        return 1
    print("\nSummary: PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main())
